# -*- coding: utf-8 -*-
"""从 config/config_v1.xlsx 导出配置（地图 + 初始选人），注入 web/index.html 的数据块。
只读 xlsx，不写回。改配置表后重新运行本脚本即可刷新网页数据。"""
import json
import re
import openpyxl

SRC = r"E:\gameDesign\rollHeroRe\config\config_v1.xlsx"
HTML = r"E:\gameDesign\rollHeroRe\web\index.html"


def load_data():
    wb = openpyxl.load_workbook(SRC, data_only=True)

    # configValue: 通用 key->content
    cfg = {}
    for row in wb["configValue"].iter_rows(min_row=2, values_only=True):
        if row and row[0]:
            cfg[str(row[0])] = row[2]
    init_tier = cfg.get("initMapTier") or 1
    team_size = int(cfg.get("initTeamSize") or 3)
    team_list = [t.strip() for t in str(cfg.get("initTeamList") or "").split(",") if t.strip()]

    # battleMap: configValue.battleMap 标准 JSON（敌方站位区块）
    # 坐标 [x,y]：x=行（1=靠近我方一行，向远处递增），y=列（1=左侧）
    battle_map = json.loads(cfg.get("battleMap") or "{}")

    # battleScene: sceneTag -> [{m,x,y}]（敌军怪物与站位）
    scene_rows = {}
    for row in wb["battleScene"].iter_rows(min_row=2, values_only=True):
        if not row or not row[1]:
            continue
        scene_rows.setdefault(str(row[1]), []).append(
            {"m": row[2], "x": row[3], "y": row[4]})
    need_monsters = {r["m"] for rows in scene_rows.values() for r in rows}

    # monster: tag -> 展示数据（只导 battleScene 引用的）
    monsters = {}
    for row in wb["monster"].iter_rows(min_row=2, values_only=True):
        if not row or not row[0] or row[0] not in need_monsters:
            continue
        monsters[row[0]] = {
            "name": row[1], "baseHp": row[4], "baseShield": row[5] or 0,
            "skills": [s.strip() for s in str(row[6] or "").split(",")
                       if s.strip() and s.strip() != "null"],
        }

    # hero: tag -> 基础数据（只导 initTeamList 里的）
    heroes = {}
    for row in wb["hero"].iter_rows(min_row=2, values_only=True):
        if not row or not row[0] or row[0] not in team_list:
            continue
        slots = [[row[4 + i], row[10 + i]] for i in range(6)]
        skills = [s.strip() for s in str(row[18] or "").split(",") if s.strip()]
        heroes[row[0]] = {
            "name": row[1], "cls": row[2], "baseHp": row[16],
            "baseShield": row[17] or 0, "slots": slots, "skills": skills,
        }

    # skill: tag -> 卡面展示 + 战斗字段（只导被英雄/怪物引用的）
    need_skills = {s for h in heroes.values() for s in h["skills"]}
    need_skills |= {s for m in monsters.values() for s in m["skills"]}
    skills = {}
    for row in wb["skill"].iter_rows(min_row=2, values_only=True):
        if not row or not row[0] or row[0] not in need_skills:
            continue
        cost = []
        for part in str(row[5] or "").split(";"):
            part = part.strip()
            if not part or part == "null":  # null = 无消耗
                continue
            k, _, v = part.partition(",")
            if k:
                cost.append([k, v])
        skills[row[0]] = {
            "name": row[2], "desc": row[3] or "", "cost": cost,
            "skillType": row[4], "useType": row[7] or "choose",
            "targetType": row[8] or "",
            "targetPosion": str(row[9] or ""),  # 受影响格子偏移 "dcol,drow;…"，相对每个最终目标展开
            "dmg": row[10] or 0,
            "dmgShieldPara": row[11] or 0,
            "shield": row[12] or 0,
        }

    # map: mapTag 向下填充，去重得到 (mapTag, tier)
    maps, seen, last = [], set(), None
    for row in wb["map"].iter_rows(min_row=2, values_only=True):
        if not row:
            continue
        tag, tier = row[0], row[1]
        if tag:
            last = tag
        if last and tier and (last, tier) not in seen:
            seen.add((last, tier))
            maps.append({"mapTag": last, "tier": tier})

    # mapPoint: A=tag B=mapTag C=pointNum D=pointType E=battleContent F-I=rewards J=extra(shop)
    points = {}
    for row in wb["mapPoint"].iter_rows(min_row=2, values_only=True):
        if not row or not isinstance(row[2], int) or not row[1]:
            continue
        _, map_tag, num, ptype, battle = row[0], row[1], row[2], row[3], row[4]
        rewards = []
        for cell in row[5:9]:
            if cell:
                k, _, v = str(cell).partition(",")
                rewards.append([k, v])
        extra = row[9] if len(row) > 9 else None
        points.setdefault(map_tag, {})[str(num)] = {
            "type": ptype,
            "battle": battle,
            "rewards": rewards,
            "extra": extra,
        }

    return {
        "initMapTier": init_tier, "maps": maps, "points": points,
        "initTeamSize": team_size, "initTeamList": team_list,
        "heroes": heroes, "skills": skills,
        "battleMap": battle_map, "battleScene": scene_rows, "monsters": monsters,
    }


def inject(data):
    with open(HTML, "r", encoding="utf-8") as f:
        html = f.read()
    js = json.dumps(data, ensure_ascii=False, separators=(",", ":"))
    new_html, n = re.subn(
        r"/\*CONFIG-DATA-START\*/.*?/\*CONFIG-DATA-END\*/",
        "/*CONFIG-DATA-START*/\nconst GAME_DATA = %s;\n/*CONFIG-DATA-END*/" % js,
        html, count=1, flags=re.S,
    )
    if n != 1:
        raise SystemExit("未找到 CONFIG-DATA 标记块")
    with open(HTML, "w", encoding="utf-8", newline="\n") as f:
        f.write(new_html)


if __name__ == "__main__":
    data = load_data()
    inject(data)
    print(f"OK: initMapTier={data['initMapTier']}, initTeamSize={data['initTeamSize']}, "
          f"initTeamList={data['initTeamList']}, heroes={len(data['heroes'])}, skills={len(data['skills'])}, "
          f"maps={len(data['maps'])}, points={sum(len(v) for v in data['points'].values())}, "
          f"battleMap={list(data['battleMap'].keys())}, "
          f"battleScene={ {k: len(v) for k, v in data['battleScene'].items()} }, monsters={len(data['monsters'])}")
    for tag, pts in data["points"].items():
        nums = sorted(int(k) for k in pts)
        missing = [n for n in range(1, max(nums) + 1) if n not in nums]
        print(f"  {tag}: {len(pts)} 格, max={max(nums)}, 不存在={missing}")
